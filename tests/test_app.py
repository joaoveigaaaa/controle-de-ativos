from copy import deepcopy
from io import BytesIO
import unittest
from unittest.mock import MagicMock
import mysql.connector
from openpyxl import Workbook, load_workbook
from app import create_app
from inventory_api import FIELDS

CONFIG = dict(DB_HOST='127.0.0.1', DB_NAME='patrimonio', DB_USER='consulta', DB_PASSWORD='test')


class InventoryTests(unittest.TestCase):
    def setUp(self):
        self.records = []
        self.connection = MagicMock()
        self.cursor = self.connection.cursor.return_value
        self.columns = [dict(Field=('numero' if f == 'numero_serie' else 'empresa' if f == 'sede' else f), Null='YES', Default=None, Extra='') for f in FIELDS]
        self.cursor.execute.side_effect = self.execute
        self.connection.start_transaction.side_effect = self.begin
        self.connection.rollback.side_effect = self.rollback
        self.connect = MagicMock(return_value=self.connection)
        self.client = create_app(CONFIG, self.connect).test_client()

    def begin(self):
        self.snapshot = deepcopy(self.records)

    def rollback(self):
        self.records[:] = self.snapshot

    def execute(self, sql, args=()):
        if sql.startswith('SHOW COLUMNS'):
            rows = self.columns
        elif sql.startswith('SELECT'):
            rows = [r for r in self.records if r['numero'] == args[0]] if args else self.records
        elif sql.startswith('INSERT'):
            columns = [x.strip().strip('`') for x in sql.split('(', 1)[1].split(')', 1)[0].split(',')]
            self.records.append(dict(zip(columns, args)))
            rows = []
        elif sql.startswith('DELETE'):
            before = len(self.records)
            self.records[:] = [r for r in self.records if r['numero'] != args[0]]
            self.cursor.rowcount = before - len(self.records)
            rows = []
        else:
            raise AssertionError(sql)
        self.cursor.fetchall.return_value = deepcopy(rows)

    def add(self, serial='00001', **values):
        return self.client.post('/api/equipamentos', json=dict(ativo='Notebook', numero_serie=serial, sede='Matriz', **values))

    def excel(self, rows, action='importar'):
        book = Workbook()
        for row in rows:
            book.active.append(row)
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        return self.client.post('/api/excel', data={'arquivo': (stream, 'test.xlsx'), 'acao': action})

    def test_register_lookup_and_inventory_contract(self):
        self.assertEqual(self.add().status_code, 201)
        self.assertEqual(self.records[0]['numero'], '00001')
        self.assertEqual(self.records[0]['empresa'], 'Matriz')
        lookup = self.client.get('/api/consulta?serie=00001').json['registros'][0]
        self.assertEqual(lookup['sede'], 'Matriz')
        self.assertEqual(lookup['numero_serie'], '00001')
        self.assertEqual(len(self.client.get('/api/equipamentos').json['registros']), 1)

    def test_connection_check_is_read_only(self):
        result = self.client.get('/api/conexao')
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json['tabela'], 'ativos')
        self.assertEqual(result.json['coluna_serie'], 'numero')
        self.assertTrue(all(call.args[0].startswith('SHOW COLUMNS') for call in self.cursor.execute.call_args_list))
        self.connection.commit.assert_not_called()
        self.assertEqual(self.records, [])

    def test_company_schema_department_and_responsible(self):
        for column in self.columns:
            if column['Field'] == 'departamento': column['Field'] = 'dp'
            if column['Field'] == 'responsavel': column['Field'] = 'nome_responsavel'
        result = self.add('AB09XZ001', departamento='TI', responsavel='Maria')
        self.assertEqual(result.status_code, 201, result.json)
        self.assertEqual(self.records[0]['dp'], 'TI')
        self.assertEqual(self.records[0]['nome_responsavel'], 'Maria')
        item = self.client.get('/api/consulta?serie=AB09XZ001').json['registros'][0]
        self.assertEqual(item['numero_serie'], 'AB09XZ001')
        self.assertEqual(item['departamento'], 'TI')

    def test_duplicate_rejected(self):
        self.add()
        self.assertEqual(self.add().status_code, 409)
        self.assertEqual(len(self.records), 1)

    def test_delete_only_selected_equipment_and_missing(self):
        self.add('A001')
        self.add('B002')
        response = self.client.delete('/api/equipamentos', json={'numero_serie':'A001'})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual([r['numero'] for r in self.records], ['B002'])
        self.assertEqual(self.client.delete('/api/equipamentos', json={'numero_serie':'A001'}).status_code, 404)

    def test_delete_ambiguous_series_is_rejected(self):
        self.records[:] = [dict(numero='X'), dict(numero='X')]
        response = self.client.delete('/api/equipamentos', json={'numero_serie':'X'})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(len(self.records), 2)

    def test_delete_cross_origin_is_rejected(self):
        response = self.client.delete('/api/equipamentos', json={'numero_serie':'A'}, headers={'Origin':'https://other.example'})
        self.assertEqual(response.status_code, 403)
        self.connect.assert_not_called()

    def test_delete_failure_rolls_back(self):
        self.add('AB01')
        original = self.execute
        def fail_delete(sql, args=()):
            original(sql, args)
            if sql.startswith('DELETE'):
                raise mysql.connector.Error(errno=1451)
        self.cursor.execute.side_effect = fail_delete
        result = self.client.delete('/api/equipamentos', json={'numero_serie':'AB01'})
        self.assertEqual(result.status_code, 409)
        self.assertEqual(self.records[0]['numero'], 'AB01')

    def test_logo_is_served_from_static(self):
        self.assertIn('src="/static/logo_luxafit.png"', self.client.get('/').get_data(as_text=True))
        response = self.client.get('/static/logo_luxafit.png')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, 'image/png')
        response.close()

    def test_absent_serial_and_sql_parameter(self):
        result = self.client.get('/api/consulta', query_string={'serie': "' OR 1=1 --"})
        self.assertEqual(result.json['registros'], [])
        self.assertEqual(self.cursor.execute.call_args.args[1], ("' OR 1=1 --",))

    def test_invalid_date_and_missing_fields(self):
        self.assertEqual(self.add(data_garantia='2026-02-30').status_code, 400)
        self.assertEqual(self.add(serial=' ').status_code, 400)
        self.assertEqual(self.records, [])

    def test_preview_import_duplicates_and_zeroes(self):
        rows = [['Ativo', 'Número de série', 'Matrícula'], ['PC', '00123', '00045'], ['PC', '00123', '00045']]
        result = self.excel(rows, 'visualizar')
        self.assertEqual(result.json['quantidade'], 2)
        self.assertEqual(self.records, [])
        self.assertEqual(self.excel(rows).json, {'importados': 1, 'ignorados': 1})
        self.assertEqual(self.records[0]['matricula'], '00045')

    def test_invalid_import_writes_nothing(self):
        response = self.excel([['Ativo', 'Série'], ['PC', 'A'], ['', 'B']])
        self.assertEqual(response.status_code, 400)
        self.assertIn('Linha 3', response.json['erro'])
        self.assertEqual(self.records, [])

    def test_import_finds_headers_after_title_and_accepts_database_names(self):
        rows = [['Inventário da empresa'], [], ['Equipamento', 'numero', 'DP', 'Nome Responsável'], ['PC', '00AB19', 'TI', 'Ana']]
        response = self.excel(rows)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(self.records[0]['numero'], '00AB19')
        self.assertEqual(self.records[0]['departamento'], 'TI')

    def test_sales_workbook_has_specific_error_and_no_database_access(self):
        response = self.excel([['VENDAS'], ['DATA', 'CLIENTE', 'VALOR', 'DATA PGTO', 'FORMA PGTO']], 'visualizar')
        self.assertEqual(response.status_code, 400)
        self.assertIn('planilha é de vendas', response.json['erro'])
        self.connect.assert_not_called()

    def test_duplicate_aliases_identify_the_column(self):
        response = self.excel([['Ativo', 'Serial', 'Número de série'], ['PC', 'A', 'B']])
        self.assertEqual(response.status_code, 400)
        self.assertIn('mais de uma coluna para Número de série', response.json['erro'])
        self.connect.assert_not_called()

    def test_missing_serial_header_is_explicit(self):
        response = self.excel([['Equipamento', 'Categoria'], ['PC', 'Computadores']])
        self.assertEqual(response.status_code, 400)
        self.assertIn('colunas obrigatórias: Número de série.', response.json['erro'])
        self.connect.assert_not_called()

    def test_db_failure_rolls_back_import(self):
        original = self.execute
        def fail_second(sql, args=()):
            if sql.startswith('INSERT') and len(self.records) == 1:
                raise mysql.connector.Error(errno=1406)
            return original(sql, args)
        self.cursor.execute.side_effect = fail_second
        response = self.excel([['Ativo', 'Série'], ['PC', 'A'], ['PC', 'B']])
        self.assertEqual(response.status_code, 503)
        self.assertEqual(self.records, [])

    def test_missing_column_rejects_instead_of_losing_data(self):
        self.columns = [c for c in self.columns if c['Field'] != 'empresa']
        response = self.add()
        self.assertEqual(response.status_code, 400)
        self.assertIn('Sede', response.json['erro'])
        self.assertEqual(self.records, [])

    def test_export_preserves_text_and_model_downloads(self):
        self.add(responsavel='=1+1')
        response = self.client.get('/api/excel/exportar')
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book.active['B2'].value, '00001')
        self.assertEqual(book.active['H2'].data_type, 's')
        book.close()
        self.assertEqual(self.client.get('/api/excel/modelo').status_code, 200)

    def test_export_deduplicates_series_and_keeps_all_columns(self):
        self.columns += [dict(Field='id', Null='NO', Default=None, Extra='auto_increment'), dict(Field='modelo', Null='YES', Default=None, Extra='')]
        self.records[:] = [
            dict(id=1, ativo='Notebook', numero=' ab001 ', empresa='Matriz', responsavel='Ana', modelo='M1'),
            dict(id=2, ativo='Notebook', numero='AB001', empresa='Filial', responsavel='', modelo='M1'),
            dict(id=3, ativo='Monitor', numero='00002', empresa='Matriz', modelo='M2'),
        ]
        response = self.client.get('/api/excel/exportar')
        self.assertEqual(response.headers['X-Export-Registros'], '2')
        self.assertEqual(response.headers['X-Export-Repetidos'], '1')
        book = load_workbook(BytesIO(response.data))
        sheet = book['Equipamentos']
        self.assertEqual(sheet.max_row, 3)
        headers = [c.value for c in sheet[1]]
        self.assertIn('modelo', headers)
        self.assertIn('id', headers)
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        chosen = next(row for row in rows if row[1] == 'AB001')
        self.assertEqual(chosen[7], 'Ana')
        self.assertEqual(chosen[8], 'Filial')
        self.assertEqual(chosen[headers.index('modelo')], 'M1')
        self.assertIn('Divergências', book.sheetnames)
        book.close()

    def test_export_does_not_merge_distinct_items_without_serial(self):
        self.records[:] = [dict(ativo='PC', numero='', empresa='Matriz'), dict(ativo='Monitor', numero='', empresa='Matriz')]
        response = self.client.get('/api/excel/exportar')
        self.assertEqual(response.headers['X-Export-Registros'], '2')

    def test_existing_html_sections_and_camera(self):
        html = self.client.get('/').get_data(as_text=True)
        for expected in ('Cadastrar equipamento', 'Importar inventário', 'Inventário', 'Iniciar câmera', 'id="serie-consulta"'):
            self.assertIn(expected, html)
        self.assertNotIn('id="foto"', html)
        self.assertIn('id="planilha"', html)

    def test_missing_config_and_db_error(self):
        response = create_app({}, self.connect).test_client().get('/api/equipamentos')
        self.assertEqual(response.status_code, 400)
        self.connect.assert_not_called()
        self.connect.side_effect = mysql.connector.Error(msg='secret-password', errno=1045)
        response = self.client.get('/api/equipamentos')
        self.assertEqual(response.status_code, 503)
        self.assertNotIn('secret-password', response.get_data(as_text=True))


if __name__ == '__main__':
    unittest.main()
